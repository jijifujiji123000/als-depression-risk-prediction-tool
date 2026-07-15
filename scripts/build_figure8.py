from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
REPORTS = ROOT / "reports"

PANELS = [
    ("A", ASSETS / "figure8_panel_home.png"),
    ("B", ASSETS / "figure8_panel_input.png"),
    ("C", ASSETS / "figure8_panel_result.png"),
    ("D", ASSETS / "figure8_panel_explanation.png"),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def fit_panel(path: Path, target_size: tuple[int, int]) -> Image.Image:
    image = Image.open(path).convert("RGB")
    # Remove the Streamlit toolbar while retaining all research-interface content.
    image = image.crop((0, 55, image.width, image.height))
    target_w, target_h = target_size
    ratio = min(target_w / image.width, target_h / image.height)
    resized = image.resize(
        (round(image.width * ratio), round(image.height * ratio)),
        Image.Resampling.LANCZOS,
    )
    frame = Image.new("RGB", target_size, "white")
    x = (target_w - resized.width) // 2
    y = (target_h - resized.height) // 2
    frame.paste(resized, (x, y))
    return frame


def build_figure() -> None:
    width, height = 3400, 2100
    margin, gap, header = 90, 70, 165
    panel_w = (width - 2 * margin - gap) // 2
    panel_h = (height - header - margin - gap) // 2
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)

    title = "Web-based ALS depressive-symptom risk prediction tool"
    title_font = font(64, bold=True)
    bbox = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 48), title, fill="#123B5D", font=title_font)
    draw.line((margin, 132, width - margin, 132), fill="#2A7F78", width=5)

    label_font = font(54, bold=True)
    for index, (label, path) in enumerate(PANELS):
        row, col = divmod(index, 2)
        x = margin + col * (panel_w + gap)
        y = header + row * (panel_h + gap)
        panel = fit_panel(path, (panel_w, panel_h))
        canvas.paste(panel, (x, y))
        draw.rounded_rectangle(
            (x, y, x + panel_w, y + panel_h),
            radius=8,
            outline="#B8C8D3",
            width=3,
        )
        draw.ellipse((x + 18, y + 18, x + 92, y + 92), fill="#123B5D")
        label_bbox = draw.textbbox((0, 0), label, font=label_font)
        label_x = x + 55 - (label_bbox[2] - label_bbox[0]) / 2
        label_y = y + 48 - (label_bbox[3] - label_bbox[1]) / 2 - label_bbox[1]
        draw.text((label_x, label_y), label, fill="white", font=label_font)

    png = REPORTS / "Figure_8_Web_Tool.png"
    pdf = REPORTS / "Figure_8_Web_Tool.pdf"
    canvas.save(png, dpi=(300, 300), optimize=True)
    canvas.save(pdf, resolution=300.0)


def build_legend() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Figure 8. Web-based ALS depressive-symptom risk prediction tool")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    legend = (
        "(A) Landing page and selection of the 3- or 6-month prediction horizon. "
        "(B) Entry of baseline clinical, scale, and laboratory variables; NLR and "
        "log1p(NLR) are calculated automatically from neutrophil and lymphocyte counts. "
        "(C) Patient-level predicted probability, risk category, prespecified threshold, "
        "and the number of locked imputation-specific models. The final probability is "
        "the arithmetic mean of predictions from 20 models. (D) Patient-level model "
        "explanation. Orange bars increase and blue bars decrease the model output; these "
        "contributions describe model behavior and should not be interpreted as causal effects. "
        "The interface does not collect identifying information or persist user-entered data."
    )
    paragraph = document.add_paragraph(legend)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(8)

    note = document.add_paragraph()
    note.add_run("Abbreviations: ").bold = True
    note.add_run(
        "ALS, amyotrophic lateral sclerosis; ALSFRS-R, Revised Amyotrophic Lateral "
        "Sclerosis Functional Rating Scale; FSS, Fatigue Severity Scale; NLR, "
        "neutrophil-to-lymphocyte ratio; PSQI, Pittsburgh Sleep Quality Index."
    )
    document.save(REPORTS / "Figure_8_Web_Tool_Legend.docx")


if __name__ == "__main__":
    REPORTS.mkdir(parents=True, exist_ok=True)
    build_figure()
    build_legend()
